"""
pytest-bdd tests for the XLSX importer.

Run with:
    python3 -m pytest tests/test_xlsx_importer_bdd.py -q

Nothing here needs a display. Converted from test_xlsx_importer.py -
every case carried over.
"""
from datetime import datetime
from types import SimpleNamespace

import pytest
from pytest_bdd import given, parsers, scenarios, then, when

from gantt_app.utils.xlsx_importer import (
    XLSXImporter, import_xlsx_file, OPENPYXL_AVAILABLE
)

pytestmark = [
    pytest.mark.xlsx_importer,
    pytest.mark.skipif(not OPENPYXL_AVAILABLE,
                       reason="openpyxl is not installed"),
]

if OPENPYXL_AVAILABLE:
    import openpyxl

scenarios("features/xlsx_import.feature")


#: A plan shaped like a hand-built spreadsheet: title block, then the table.
SAMPLE_ROWS = [
    ('Implementation Plan',),
    (),
    ('ID', 'Phase', 'Task', 'Pred.', 'Duration (wd)', 'Start', 'End',
     'Status'),
    (1, 'Phase One', 'Kick-off', '–', 5,
     datetime(2024, 1, 1), datetime(2024, 1, 5), 'Ongoing'),
    (2, 'Phase One', 'Analysis', '1', 5,
     datetime(2024, 1, 8), datetime(2024, 1, 12), 'Not started'),
    (3, 'Phase Two', 'Build', '1;2', 10,
     datetime(2024, 1, 15), datetime(2024, 1, 26), 'Done'),
]


# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------

def build_workbook(rows, sheet_title="Plan", path=None):
    """Write row tuples to a temporary .xlsx file and return its path."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path


def coerce(cell):
    """Turn a feature-table cell back into the value a sheet would hold."""
    if cell is None or cell == '':
        return None
    try:
        return int(cell)
    except ValueError:
        pass
    try:
        return float(cell)
    except ValueError:
        pass
    try:
        return datetime.fromisoformat(cell)
    except ValueError:
        pass
    return cell


def rows_from(datatable):
    """A datatable's rows, coerced, with all-empty rows kept as ()."""
    rows = []
    for row in datatable:
        values = [coerce(cell) for cell in row]
        rows.append(()) if all(v is None for v in values) \
            else rows.append(tuple(values))
    return rows


def task_named(project, name):
    return next(t for t in project.tasks if t.name == name)


# ------------------------------------------------------------------
# WHEN - importing
# ------------------------------------------------------------------

@when("the sample worksheet is imported", target_fixture="ctx")
def the_sample_worksheet_is_imported(tmp_path):
    path = tmp_path / "sample.xlsx"
    build_workbook(SAMPLE_ROWS, path=str(path))
    return SimpleNamespace(project=XLSXImporter().import_xlsx(str(path)))


@when("the sample worksheet is imported without phase grouping",
      target_fixture="ctx")
def the_sample_worksheet_is_imported_flat(tmp_path):
    path = tmp_path / "sample.xlsx"
    build_workbook(SAMPLE_ROWS, path=str(path))
    importer = XLSXImporter(group_by_phase=False)
    return SimpleNamespace(project=importer.import_xlsx(str(path)))


@when("the sample worksheet is imported through the convenience function",
      target_fixture="ctx")
def the_sample_worksheet_via_the_function(tmp_path):
    path = tmp_path / "sample.xlsx"
    build_workbook(SAMPLE_ROWS, path=str(path))
    return SimpleNamespace(project=import_xlsx_file(str(path)))


@when("this worksheet is imported", target_fixture="ctx")
def this_worksheet_is_imported(tmp_path, datatable):
    path = tmp_path / "sheet.xlsx"
    build_workbook(rows_from(datatable), path=str(path))
    return SimpleNamespace(project=XLSXImporter().import_xlsx(str(path)))


@when(parsers.parse('this worksheet is imported on a sheet named "{title}"'),
      target_fixture="ctx")
def this_worksheet_is_imported_on_a_named_sheet(tmp_path, datatable, title):
    path = tmp_path / "sheet.xlsx"
    build_workbook(rows_from(datatable), sheet_title=title, path=str(path))
    return SimpleNamespace(project=XLSXImporter().import_xlsx(str(path)))


@when(parsers.parse('a workbook is imported whose cover sheet holds notes '
                    'and whose "{data_sheet}" sheet holds a table'),
      target_fixture="ctx")
def a_workbook_with_a_cover_sheet(tmp_path, data_sheet):
    workbook = openpyxl.Workbook()
    cover = workbook.active
    cover.title = "Cover"
    cover.append(["Some notes"])
    cover.append(["No table here"])

    data = workbook.create_sheet(data_sheet)
    for row in [('ID', 'Task', 'Start', 'Duration'),
                (1, 'Real task', datetime(2024, 1, 1), 3)]:
        data.append(list(row))

    path = tmp_path / "covered.xlsx"
    workbook.save(str(path))
    return SimpleNamespace(project=XLSXImporter().import_xlsx(str(path)))


@when("a nonexistent XLSX file is imported", target_fixture="ctx")
def a_nonexistent_xlsx_file(tmp_path):
    missing = tmp_path / "no-such-plan.xlsx"
    return SimpleNamespace(project=XLSXImporter().import_xlsx(str(missing)))


# ------------------------------------------------------------------
# GIVEN - the round trip
# ------------------------------------------------------------------

@given("the sample worksheet has been imported and re-exported",
       target_fixture="ctx")
def the_sample_worksheet_round_tripped(tmp_path):
    from gantt_app.utils.xlsx_exporter import export_project_to_xlsx

    source = tmp_path / "source.xlsx"
    build_workbook(SAMPLE_ROWS, path=str(source))

    original = import_xlsx_file(str(source))
    # Settled before exporting, which is what the application does the
    # moment a plan is loaded. A re-imported plan comes back settled -
    # the sheet's dates are formulas, so the scheduler is what places the
    # rows - and comparing a settled plan against an unsettled one would
    # differ on the summary rows, whose progress is rolled up.
    original.reschedule()

    exported = tmp_path / "exported.xlsx"
    assert export_project_to_xlsx(original, str(exported))
    reimported = import_xlsx_file(str(exported))
    return SimpleNamespace(original=original, reimported=reimported)


# ------------------------------------------------------------------
# GIVEN - logging
# ------------------------------------------------------------------

@given("the log is being watched")
def the_log_is_being_watched():
    from gantt_app.utils.log import setup_logging, reset_logging
    reset_logging()
    setup_logging(to_file=False, to_stderr=False)


# ------------------------------------------------------------------
# THEN - imported plans
# ------------------------------------------------------------------

@then(parsers.parse('the imported project is named "{name}"'))
def the_project_is_named(ctx, name):
    assert ctx.project is not None
    assert ctx.project.name == name


@then(parsers.parse('the imported project is not named "{name}"'))
def the_project_is_not_named(ctx, name):
    assert ctx.project.name != name


@then(parsers.parse('the imported task names include "{a}", "{b}" and '
                    '"{c}"'))
def the_task_names_include(ctx, a, b, c):
    names = [t.name for t in ctx.project.tasks]
    for name in (a, b, c):
        assert name in names


@then(parsers.parse('it holds {count:d} tasks'))
def it_holds_tasks(ctx, count):
    assert ctx.project is not None
    assert len(ctx.project.tasks) == count


@then(parsers.parse('it holds {count:d} tasks, all at the top level'))
def it_holds_flat_tasks(ctx, count):
    assert len(ctx.project.tasks) == count
    assert all(t.parent_task_id is None for t in ctx.project.tasks)


@then("nothing comes back")
def nothing_comes_back(ctx):
    assert ctx.project is None


@then(parsers.parse('the imported task "{task_id}" starts on "{day}"'))
def the_task_starts_on(ctx, task_id, day):
    task = ctx.project.get_task_by_id(task_id)
    assert task.start_date == datetime.fromisoformat(day)


@then(parsers.parse('the imported task "{task_id}" ends on "{day}"'))
def the_task_ends_on(ctx, task_id, day):
    task = ctx.project.get_task_by_id(task_id)
    assert task.end_date == datetime.fromisoformat(day)


@then(parsers.parse('the imported task "{task_id}" lasts {days:d} days'))
def the_task_lasts(ctx, task_id, days):
    assert ctx.project.get_task_by_id(task_id).duration_days == days


@then(parsers.parse('the imported task "{task_id}" waits on "{a}" and '
                    '"{b}"'))
def the_task_waits_on_two(ctx, task_id, a, b):
    task = ctx.project.get_task_by_id(task_id)
    assert sorted(task.dependency_ids) == sorted([a, b])


@then(parsers.parse('the imported task "{task_id}" waits on nothing'))
def the_task_waits_on_nothing(ctx, task_id):
    task = ctx.project.get_task_by_id(task_id)
    assert task.dependency_ids == [] or list(task.dependencies) == []


@then(parsers.parse('the imported task "{task_id}" is {percent:d} percent '
                    'done'))
def the_task_is_percent_done(ctx, task_id, percent):
    assert ctx.project.get_task_by_id(task_id).progress == percent


@then(parsers.parse('the imported task "{task_id}" is named "{name}"'))
def the_task_is_named(ctx, task_id, name):
    assert ctx.project.get_task_by_id(task_id).name == name


@then(parsers.parse('the imported task "{task_id}" is a milestone with no '
                    'end date'))
def the_task_is_a_milestone(ctx, task_id):
    task = ctx.project.get_task_by_id(task_id)
    assert task.is_milestone
    assert task.end_date is None


@then(parsers.parse('the imported task "{task_id}" has parent "{parent}" '
                    'and type "{task_type}"'))
def the_task_has_parent_and_type(ctx, task_id, parent, task_type):
    task = ctx.project.get_task_by_id(task_id)
    assert task.parent_task_id == parent
    assert task.task_type == task_type


@then(parsers.parse('the root tasks are "{a}" and "{b}"'))
def the_root_tasks_are(ctx, a, b):
    assert [t.name for t in ctx.project.get_root_tasks()] == [a, b]


@then(parsers.parse('"{phase}" holds "{a}" and "{b}" as subtasks'))
def the_phase_holds_subtasks(ctx, phase, a, b):
    parent = next(t for t in ctx.project.get_root_tasks()
                  if t.name == phase)
    subtasks = ctx.project.get_subtasks(parent.id)
    assert [t.name for t in subtasks] == [a, b]
    for subtask in subtasks:
        assert subtask.task_type == "Subtask"


@then(parsers.parse('the imported "{name}" starts on "{start}" and ends on '
                    '"{end}"'))
def the_task_spans(ctx, name, start, end):
    task = task_named(ctx.project, name)
    assert task.start_date == datetime.fromisoformat(start)
    assert task.end_date == datetime.fromisoformat(end)


@then(parsers.parse('the imported "{name}" waits on the imported "{other}"'))
def the_task_waits_on_the_task(ctx, name, other):
    task = task_named(ctx.project, name)
    predecessor = task_named(ctx.project, other)
    assert task.dependency_ids == [predecessor.id]


@then("every imported task id is unique")
def every_task_id_is_unique(ctx):
    ids = [t.id for t in ctx.project.tasks]
    assert len(ids) == len(set(ids)), f"duplicate ids: {ids}"


@then(parsers.parse('the log mentions "{text}"'))
def the_log_mentions(text):
    from gantt_app.utils.log import get_log_text, reset_logging
    try:
        assert text in get_log_text()
    finally:
        reset_logging()


# ------------------------------------------------------------------
# THEN - the round trip
# ------------------------------------------------------------------

@then("the re-import produced a project")
def the_reimport_produced_a_project(ctx):
    assert ctx.reimported is not None


@then("the re-import holds as many tasks as the original")
def the_reimport_holds_as_many(ctx):
    assert len(ctx.reimported.tasks) == len(ctx.original.tasks)


@then("the re-import is named after the original")
def the_reimport_is_named(ctx):
    assert ctx.reimported.name == ctx.original.name


@then("every re-imported task matches the original's dates, progress and "
      "milestone flag")
def every_task_matches(ctx):
    original = {t.name: t for t in ctx.original.tasks}
    reimported = {t.name: t for t in ctx.reimported.tasks}
    assert set(original) == set(reimported)
    for name, task in original.items():
        other = reimported[name]
        assert task.start_date == other.start_date, name
        assert task.end_date == other.end_date, name
        assert task.progress == other.progress, name
        assert task.is_milestone == other.is_milestone, name


@then("every re-imported task waits on the same named tasks")
def the_dependencies_survive(ctx):
    original = {t.name: t for t in ctx.original.tasks}
    reimported = {t.name: t for t in ctx.reimported.tasks}
    for name, task in original.items():
        expected = sorted(ctx.original.get_task_by_id(d).name
                          for d in task.dependency_ids)
        actual = sorted(ctx.reimported.get_task_by_id(d).name
                        for d in reimported[name].dependency_ids)
        assert expected == actual, name


@then("every re-imported task has the same named parent")
def the_hierarchy_survives(ctx):
    original = {t.name: t for t in ctx.original.tasks}
    reimported = {t.name: t for t in ctx.reimported.tasks}
    for name, task in original.items():
        expected = (ctx.original.get_task_by_id(task.parent_task_id).name
                    if task.parent_task_id else None)
        other = reimported[name]
        actual = (ctx.reimported.get_task_by_id(other.parent_task_id).name
                  if other.parent_task_id else None)
        assert expected == actual, name


@then("the re-import has as many root tasks as the original")
def the_roots_match(ctx):
    assert len(ctx.reimported.get_root_tasks()) == \
        len(ctx.original.get_root_tasks())


# ------------------------------------------------------------------
# THEN - value coercions
# ------------------------------------------------------------------

@then(parsers.parse('"{raw}" normalises to "{clean}"'))
def the_header_normalises(raw, clean):
    assert XLSXImporter()._normalise_header(raw) == clean


@then(parsers.parse('"{cell}" splits into "{a}" and "{b}"'))
def the_cell_splits_into_two(cell, a, b):
    assert XLSXImporter()._split_dependencies(cell) == [a, b]


@then(parsers.parse('"{cell}" splits into "{a}", "{b}" and "{c}"'))
def the_cell_splits_into_three(cell, a, b, c):
    assert XLSXImporter()._split_dependencies(cell) == [a, b, c]


@then(parsers.parse('"{cell}" splits into nothing'))
def the_cell_splits_into_nothing(cell):
    assert XLSXImporter()._split_dependencies(cell) == []


@then("an empty cell splits into nothing")
def an_empty_cell_splits_into_nothing():
    assert XLSXImporter()._split_dependencies(None) == []


@then(parsers.parse('the number {num:d} splits into "{name}"'))
def the_number_splits_into(num, name):
    assert XLSXImporter()._split_dependencies(num) == [name]


@then(parsers.parse('"{cell}" splits into "{a}"'))
def the_cell_splits_into_one(cell, a):
    assert XLSXImporter()._split_dependencies(cell) == [a]


@then(parsers.parse('"{cell}" splits into itself'))
def the_cell_splits_into_itself(cell):
    assert XLSXImporter()._split_dependencies(cell) == [cell]


@then(parsers.parse('the cell "{text}" parses to "{day}"'))
def the_cell_parses_to(text, day):
    value = coerce(text)
    assert XLSXImporter()._parse_cell_date(value) == \
        datetime.fromisoformat(day)


@then(parsers.parse('the cell serial "{serial}" parses to "{day}"'))
def the_serial_parses_to(serial, day):
    assert XLSXImporter()._parse_cell_date(int(serial)) == \
        datetime.fromisoformat(day)


@then(parsers.parse('the cell "{text}" parses to nothing'))
def the_cell_parses_to_nothing(text):
    assert XLSXImporter()._parse_cell_date(text) is None


@then("an empty cell parses to nothing")
def an_empty_cell_parses_to_nothing():
    assert XLSXImporter()._parse_cell_date(None) is None


@then(parsers.parse('{days:d} working days from "{start}" end on "{end}"'))
def working_days_end_on(days, start, end):
    assert XLSXImporter()._end_date_for(
        datetime.fromisoformat(start), days, True) == \
        datetime.fromisoformat(end)


@then(parsers.parse('{days:d} calendar days from "{start}" end on "{end}"'))
def calendar_days_end_on(days, start, end):
    assert XLSXImporter()._end_date_for(
        datetime.fromisoformat(start), days, False) == \
        datetime.fromisoformat(end)


@then(parsers.parse('{days:d} working days ending "{end}" started on '
                    '"{start}"'))
def working_days_started_on(days, end, start):
    assert XLSXImporter()._start_date_for(
        datetime.fromisoformat(end), days, True) == \
        datetime.fromisoformat(start)


@then(parsers.parse('{days:d} calendar days ending "{end}" started on '
                    '"{start}"'))
def calendar_days_started_on(days, end, start):
    assert XLSXImporter()._start_date_for(
        datetime.fromisoformat(end), days, False) == \
        datetime.fromisoformat(start)


@then(parsers.parse('progress {progress} with status "{status}" reads as '
                    '{expected:d}'))
def progress_with_status_reads(progress, status, expected):
    row = {'progress': coerce(progress), 'status': status}
    assert XLSXImporter()._progress_from_row(row) == expected


@then(parsers.parse('progress {progress} reads as {expected:d}'))
def progress_reads(progress, expected):
    row = {'progress': coerce(progress)}
    assert XLSXImporter()._progress_from_row(row) == expected


@then(parsers.parse('status "{status}" reads as {expected:d}'))
def status_reads(status, expected):
    assert XLSXImporter()._progress_from_row({'status': status}) == expected


@then(parsers.parse('no row data reads as {expected:d}'))
def no_row_data_reads(expected):
    assert XLSXImporter()._progress_from_row({}) == expected
