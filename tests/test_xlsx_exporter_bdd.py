"""
pytest-bdd tests for the XLSX export: the plan sheet the exporter writes.

Run with:
    python3 -m pytest tests/test_xlsx_exporter_bdd.py -q

Nothing here needs a display. Converted from test_xlsx_exporter.py -
every case carried over.
"""
from datetime import date, datetime, timedelta
from io import BytesIO
from types import SimpleNamespace

import pytest
from pytest_bdd import given, parsers, scenarios, then, when

from gantt_app.core.models import Project, Task
from gantt_app.core.workdaycalendar import DateOverride
from gantt_app.utils.xlsx_exporter import (
    OPENPYXL_AVAILABLE, export_project_to_xlsx, generate_xlsx_bytes,
)

pytestmark = [
    pytest.mark.xlsx_exporter,
    pytest.mark.skipif(not OPENPYXL_AVAILABLE,
                       reason="openpyxl is not installed"),
]

if OPENPYXL_AVAILABLE:
    import openpyxl

scenarios("features/xlsx_export.feature")


# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------

def build_project() -> Project:
    """A small plan starting on Monday 6 July 2026."""
    project = Project(name="Tosca Implementation")
    base = datetime(2026, 7, 6)

    project.add_task(Task(id="P1", name="1. Procurement", task_type="Phase",
                          start_date=base, end_date=base))
    project.add_task(Task(id="D1", name="Signed contract",
                          task_type="Task", parent_task_id="P1",
                          start_date=base, end_date=base))
    project.add_task(Task(id="T1", name="Business case", task_type="Subtask",
                          parent_task_id="D1", start_date=base,
                          details="Signed contract",
                          end_date=base + timedelta(days=4), progress=50))
    project.add_task(Task(id="T2", name="Procurement demand",
                          task_type="Subtask", parent_task_id="D1",
                          start_date=base, end_date=base + timedelta(days=9)))
    project.get_task_by_id("T2").add_dependency("T1", "FS", "Hard")

    project.add_task(Task(id="P2", name="2. Requirements", task_type="Phase",
                          start_date=base, end_date=base))
    project.add_task(Task(id="T3", name="URS", task_type="Subtask",
                          parent_task_id="P2", start_date=base,
                          end_date=base + timedelta(days=9), progress=100))
    project.get_task_by_id("T3").add_dependency("T2", "FS", "Hard")

    project.reschedule()
    return project


def export(ctx, tmp_path):
    """Write the project out and keep its plan sheet on the context."""
    path = tmp_path / "plan.xlsx"
    assert export_project_to_xlsx(ctx.project, str(path))
    ctx.workbook = openpyxl.load_workbook(str(path))
    ctx.sheet = ctx.workbook.worksheets[0]


def row_of(sheet, name):
    """The sheet row a named task was written to."""
    for row in range(6, sheet.max_row + 1):
        if sheet.cell(row=row, column=3).value == name:
            return row
    raise AssertionError(f"no row for {name!r}")


def column(ctx, row, letter):
    """One cell by row and column letter."""
    return ctx.sheet[f"{letter}{row}"].value


# ------------------------------------------------------------------
# GIVEN
# ------------------------------------------------------------------

@given("the Tosca plan", target_fixture="ctx")
def the_tosca_plan():
    return SimpleNamespace(project=build_project())


@given(parsers.parse('a plan holding only a phase named "{name}"'),
       target_fixture="ctx")
def a_plan_holding_only_a_phase(name):
    project = Project(name="Just a phase")
    project.add_task(Task(id="P", name=name, task_type="Phase",
                          start_date=datetime(2026, 7, 6),
                          end_date=datetime(2026, 7, 10)))
    return SimpleNamespace(project=project)


@given(parsers.parse('the Tosca plan with "{third}" following "{second}" '
                     'start-to-start'), target_fixture="ctx")
def the_tosca_plan_with_an_ss_link(third, second):
    project = build_project()
    task = project.get_task_by_id("T3")
    task.dependencies = []
    task.add_dependency("T2", "SS", "Hard")
    project.apply_dependency_constraints(task)
    project.reschedule()
    return SimpleNamespace(project=project)


@given(parsers.parse('the Tosca plan with "{second}" lagged {lag:d} days '
                     'behind "{first}"'), target_fixture="ctx")
def the_tosca_plan_with_a_lagged_link(second, lag, first):
    project = build_project()
    task = project.get_task_by_id("T2")
    task.dependencies = []
    task.add_dependency("T1", "FS", "Hard", lag=lag)
    project.apply_dependency_constraints(task)
    project.reschedule()
    return SimpleNamespace(project=project)


@given(parsers.parse('a plan with tasks starting "{early}" and "{later}"'),
       target_fixture="ctx")
def a_plan_with_two_starts(early, later):
    project = Project(name="Two starts")
    project.add_task(Task(id="A", name="Early",
                          start_date=datetime.fromisoformat(early),
                          end_date=datetime(2026, 7, 10)))
    project.add_task(Task(id="B", name="Later",
                          start_date=datetime.fromisoformat(later),
                          end_date=datetime(2026, 8, 7)))
    project.reschedule()
    return SimpleNamespace(project=project)


@given(parsers.parse('the Tosca plan with a milestone "{name}" on "{day}"'),
       target_fixture="ctx")
def the_tosca_plan_with_a_milestone(name, day):
    project = build_project()
    project.add_task(Task(id="M", name=name, task_type="Milestone",
                          parent_task_id="P2",
                          start_date=datetime.fromisoformat(day)))
    project.reschedule()
    return SimpleNamespace(project=project)


@given("the Tosca plan with Hungarian holidays", target_fixture="ctx")
def the_tosca_plan_with_holidays():
    project = build_project()
    project.set_holiday_countries(["HU"])
    return SimpleNamespace(project=project)


@given("the Tosca plan with a manual shutdown on its second working day",
       target_fixture="ctx")
def the_tosca_plan_with_a_manual_shutdown():
    project = build_project()
    shutdown = project.start_date.date() + timedelta(days=1)
    while shutdown.weekday() >= 5:
        shutdown += timedelta(days=1)
    project.set_date_overrides([
        DateOverride(shutdown, False, "Company shutdown")])
    ctx = SimpleNamespace(project=project)
    ctx.shutdown = shutdown
    return ctx


@given(parsers.parse('a plan working "{day}" as a make-up day'),
       target_fixture="ctx")
def a_plan_working_a_makeup_day(day):
    project = Project(name="Make-up")
    project.add_task(Task(id="T1", name="Over the Saturday",
                          start_date=datetime(2026, 9, 11),
                          end_date=datetime(2026, 9, 14)))
    project.reschedule()
    project.set_date_overrides([
        DateOverride(datetime.fromisoformat(day).date(), True,
                     "Make-up day")])
    return SimpleNamespace(project=project)


@given(parsers.parse('a plan resting "{day}" for a company shutdown'),
       target_fixture="ctx")
def a_plan_resting_a_shutdown(day):
    project = Project(name="Shutdown")
    project.add_task(Task(id="T1", name="Over the shutdown",
                          start_date=datetime(2026, 9, 14),
                          end_date=datetime(2026, 9, 16)))
    project.reschedule()
    project.set_date_overrides([
        DateOverride(datetime.fromisoformat(day).date(), False,
                     "Company shutdown")])
    return SimpleNamespace(project=project)


@given(parsers.parse('a plan holding one task named "{name}"'),
       target_fixture="ctx")
def a_plan_holding_one_task(name):
    project = Project(name="In memory")
    project.add_task(Task(id="A", name=name,
                          start_date=datetime(2026, 7, 6),
                          end_date=datetime(2026, 7, 10)))
    return SimpleNamespace(project=project)


# ------------------------------------------------------------------
# WHEN
# ------------------------------------------------------------------

@when("the plan is exported to a workbook")
def the_plan_is_exported_to_a_workbook(ctx, tmp_path):
    export(ctx, tmp_path)


@when("the plan is exported to workbook bytes")
def the_plan_is_exported_to_bytes(ctx):
    ctx.data = generate_xlsx_bytes(ctx.project)


# ------------------------------------------------------------------
# THEN - the sheet's fixed parts
# ------------------------------------------------------------------

@then(parsers.parse('the plan sheet is titled "{title}"'))
def the_sheet_is_titled(ctx, title):
    assert ctx.sheet.title == title


@then(parsers.parse('cell "{cell}" names "{text}"'))
def the_cell_names(ctx, cell, text):
    assert text in str(ctx.sheet[cell].value)


@then(parsers.parse('cell "{cell}" reads "{text}"'))
def the_cell_reads(ctx, cell, text):
    assert ctx.sheet[cell].value == text


@then("row 5 carries the ten plan headings")
def the_headings_are_the_plan_columns(ctx):
    headings = [ctx.sheet.cell(row=5, column=c).value for c in range(1, 11)]
    assert headings == [
        'ID', 'Phase', 'Task', 'Responsible (A)', 'Key Deliverable',
        'Pred.', 'Duration (wd)', 'Start', 'End', 'Status',
    ]


@then(parsers.parse('the freeze pane is "{cell}"'))
def the_freeze_pane_is(ctx, cell):
    assert ctx.sheet.freeze_panes == cell


@then(parsers.parse('a "{sheet}" sheet names the project "{name}"'))
def the_totals_are_on_their_own_sheet(ctx, sheet, name):
    assert sheet in ctx.workbook.sheetnames
    summary = ctx.workbook[sheet]
    labels = {summary.cell(row=r, column=1).value:
              summary.cell(row=r, column=2).value
              for r in range(1, summary.max_row + 1)}
    assert labels["Project Name:"] == name


# ------------------------------------------------------------------
# THEN - which tasks get rows
# ------------------------------------------------------------------

@then(parsers.parse('the task column lists "{a}", "{b}" and "{c}"'))
def the_task_column_lists(ctx, a, b, c):
    names = [ctx.sheet.cell(row=r, column=3).value
             for r in range(6, ctx.sheet.max_row + 1)]
    assert names == [a, b, c]


@then(parsers.parse('the task column does not list "{name}"'))
def the_task_column_omits(ctx, name):
    names = [ctx.sheet.cell(row=r, column=3).value
             for r in range(6, ctx.sheet.max_row + 1)]
    assert name not in names


@then(parsers.parse('the row for "{name}" carries phase "{phase}"'))
def the_row_carries_phase(ctx, name, phase):
    assert column(ctx, row_of(ctx.sheet, name), 'B') == phase


@then(parsers.parse('the row for "{name}" carries key deliverable "{text}"'))
def the_row_carries_key_deliverable(ctx, name, text):
    assert column(ctx, row_of(ctx.sheet, name), 'E') == text


@then(parsers.parse('the phase cells for "{a}" and "{b}" differ in colour'))
def the_phase_cells_differ(ctx, a, b):
    first = ctx.sheet[f"B{row_of(ctx.sheet, a)}"].fill.fgColor.rgb
    second = ctx.sheet[f"B{row_of(ctx.sheet, b)}"].fill.fgColor.rgb
    assert first != second


@then(parsers.parse('the first task row reads "{name}"'))
def the_first_task_row_reads(ctx, name):
    assert ctx.sheet.cell(row=6, column=3).value == name


# ------------------------------------------------------------------
# THEN - the live sheet
# ------------------------------------------------------------------

@then(parsers.parse('the row for "{name}" starts with "{formula}"'))
def the_row_starts_with(ctx, name, formula):
    assert column(ctx, row_of(ctx.sheet, name), 'H') == formula


@then(parsers.parse('the row for "{name}" starts the working day after '
                    '"{other}" finishes'))
def the_row_chains_on(ctx, name, other):
    row = row_of(ctx.sheet, other)
    assert column(ctx, row_of(ctx.sheet, name), 'H') == \
        f"=WORKDAY(I{row},1)"


@then(parsers.parse('the row for "{name}" finishes by a plain WORKDAY'))
def the_row_finishes_by_workday(ctx, name):
    row = row_of(ctx.sheet, name)
    assert column(ctx, row, 'I') == f"=WORKDAY(H{row},G{row}-1)"


@then(parsers.parse('the row for "{name}" holds the task\'s working '
                    'duration'))
def the_row_holds_the_working_duration(ctx, name):
    task = ctx.project.get_task_by_id("T1")
    assert column(ctx, row_of(ctx.sheet, name), 'G') == \
        ctx.project.working_duration(task)


@then(parsers.parse('the row for "{name}" draws its week K bar by overlap'))
def the_timeline_is_drawn_from_the_dates(ctx, name):
    row = row_of(ctx.sheet, name)
    assert column(ctx, row, 'K') == \
        f'=IF(AND(K$5<=$I{row},K$5+6>=$H{row}),"█","")'


@then(parsers.parse('cell "{first}" reads "{first_f}" and cell "{second}" '
                    'reads "{second_f}"'))
def the_weeks_run_from_the_start(ctx, first, first_f, second, second_f):
    assert ctx.sheet[first].value == first_f
    assert ctx.sheet[second].value == second_f


# ------------------------------------------------------------------
# THEN - formulas agree with the plan
# ------------------------------------------------------------------

@then(parsers.parse('the row for "{name}" starts with the task\'s own '
                    'start date'))
def the_row_starts_with_the_date(ctx, name):
    task = next(t for t in ctx.project.tasks if t.name == name)
    row = next(r for r in range(6, ctx.sheet.max_row + 1)
               if ctx.sheet.cell(row=r, column=3).value == name)
    assert ctx.sheet[f"H{row}"].value == task.start_date
    # And it is not just the project start cell in disguise
    assert task.start_date != ctx.project.start_date


@then(parsers.parse('the row for "{name}" reads its predecessor as "{tid}" '
                    'with an "{suffix}" suffix'))
def the_row_reads_its_predecessor(ctx, name, tid, suffix):
    row = next(r for r in range(6, ctx.sheet.max_row + 1)
               if ctx.sheet.cell(row=r, column=3).value == name)
    assert ctx.sheet[f"F{row}"].value == \
        f"{ctx.project.display_ids()[tid]}{suffix}"


@then(parsers.parse('the first row starts with "{formula}"'))
def the_first_row_starts_with(ctx, formula):
    assert ctx.sheet["H6"].value == formula


@then(parsers.parse('the second row starts with the date "{day}"'))
def the_second_row_starts_with_the_date(ctx, day):
    assert ctx.sheet["H7"].value == datetime.fromisoformat(day)


@then(parsers.parse('the row for "{name}" reads its predecessor as "{dash}"'))
def the_row_reads_a_dash(ctx, name, dash):
    assert column(ctx, row_of(ctx.sheet, name), 'F') == dash


# ------------------------------------------------------------------
# THEN - milestones and status
# ------------------------------------------------------------------

@then(parsers.parse('the row for "{name}" has a zero duration'))
def the_row_has_zero_duration(ctx, name):
    row = next(r for r in range(6, ctx.sheet.max_row + 1)
               if ctx.sheet.cell(row=r, column=3).value == name)
    assert ctx.sheet[f"G{row}"].value == 0


@then(parsers.parse('the row for "{name}" finishes on its own start'))
def the_row_finishes_on_its_start(ctx, name):
    row = next(r for r in range(6, ctx.sheet.max_row + 1)
               if ctx.sheet.cell(row=r, column=3).value == name)
    assert ctx.sheet[f"I{row}"].value == f"=$H{row}"


@then(parsers.parse('the row for "{name}" reads status "{status}"'))
def the_row_reads_status(ctx, name, status):
    assert column(ctx, row_of(ctx.sheet, name), 'J') == status


@then(parsers.parse('the status cell for "{name}" is filled with the '
                    'ongoing colour'))
def the_status_cell_colour(ctx, name):
    from gantt_app.utils.xlsx_exporter import STATUS_FILLS, STATUS_ONGOING
    cell = ctx.sheet[f"J{row_of(ctx.sheet, name)}"]
    assert cell.fill.fgColor.rgb == STATUS_FILLS[STATUS_ONGOING]


# ------------------------------------------------------------------
# THEN - holidays
# ------------------------------------------------------------------

@then(parsers.parse('there is no "{sheet}" sheet'))
def there_is_no_sheet(ctx, sheet):
    assert sheet not in ctx.workbook.sheetnames


@then(parsers.parse('a hidden "{sheet}" sheet exists'))
def a_hidden_sheet_exists(ctx, sheet):
    assert sheet in ctx.workbook.sheetnames
    assert ctx.workbook[sheet].sheet_state == 'hidden'


@then(parsers.parse('the row for "{name}" finishes by a WORKDAY over the '
                    'holiday sheet'))
def the_row_uses_the_holiday_sheet(ctx, name):
    row = next(r for r in range(6, ctx.sheet.max_row + 1)
               if ctx.sheet.cell(row=r, column=3).value == name)
    assert ctx.sheet[f"I{row}"].value == \
        f"=WORKDAY(H{row},G{row}-1,Holidays!$A:$A)"


@then(parsers.parse('the "{sheet}" sheet lists the shutdown date'))
def the_holiday_sheet_lists_the_shutdown(ctx, sheet):
    listed = [cell[0].value for cell in ctx.workbook[sheet].iter_rows()]
    listed = [v.date() if hasattr(v, 'date') else v for v in listed]
    assert ctx.shutdown in listed


@then(parsers.parse('the "{sheet}" sheet exists'))
def the_sheet_exists(ctx, sheet):
    assert sheet in ctx.workbook.sheetnames


@then(parsers.parse('the task ends on "{day}" and its row\'s finish is the '
                    'date itself'))
def the_task_over_a_worked_saturday(ctx, day):
    task = ctx.project.get_task_by_id("T1")
    assert task.end_date.date() == datetime.fromisoformat(day).date()
    assert ctx.sheet["I6"].value == task.end_date


@then("the first row finishes by a WORKDAY over the holiday sheet")
def the_first_row_uses_the_holiday_sheet(ctx):
    assert ctx.sheet["I6"].value == "=WORKDAY(H6,G6-1,Holidays!$A:$A)"


# ------------------------------------------------------------------
# THEN - the bytes entry point
# ------------------------------------------------------------------

@then(parsers.parse('the bytes open as a sheet whose first task row reads '
                    '"{name}"'))
def the_bytes_open_as_a_workbook(ctx, name):
    assert ctx.data is not None
    workbook = openpyxl.load_workbook(BytesIO(ctx.data))
    assert workbook.worksheets[0].cell(row=6, column=3).value == name
